"""Importación de clientes y equipos desde la exportación de Persat (E11-01, E11-02).

Todo en una transacción: en modo prueba se deshace al final, así las cantidades y los
códigos que informa son los mismos que daría la importación real. Lo ya importado se
reconoce por su referencia a Persat (`legacy_reference`) y no se vuelve a crear.
"""

import uuid
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.catalog import EquipmentType
from app.db.models.client import Client, ClientContact, ClientSegment, ClientSite, TaxCondition
from app.db.models.equipment import Equipment
from app.db.models.legacy_reference import LegacyReference
from app.db.repositories.client import ClientRepository, ClientSiteRepository
from app.db.repositories.equipment import EquipmentRepository
from app.domain import persat
from app.services.equipment import next_code

SOURCE = "persat"
CLIENT_ENTITY = "client"
EQUIPMENT_ENTITY = "equipment"

CLIENT_HEADERS = (
    "Nro. de Cliente",
    "Razón Social",
    "Descripción",
    "Tipo",
    "Dirección",
    "Nombre",
    "Email",
    "Telefono",
)
EQUIPMENT_HEADERS = (
    "Numero de Serie",
    "Nro. de Cliente",
    "UBICACION",
    "TIPO",
    "SUB TIPO",
    "MODELO",
    "MARCA",
    "POTENCIA",
    "GAS REFRIGERANTE",
    "TENSION DE ALIMENTACION",
)
NO_ADDRESS = "Sin dirección"

# Largos máximos de las columnas que se llenan con texto de Persat.
_MAX_LENGTHS = {
    "número": 20,
    "razón social": 255,
    "dirección": 255,
    "referencia": 500,
    "código": 50,
    "ubicación": 255,
}

Row = dict[str, str]


class PersatFileError(Exception):
    """El archivo no se puede leer o no tiene las columnas de la exportación de Persat."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


@dataclass(frozen=True)
class Issue:
    """Una fila del archivo con algo que informar. Sin nombres ni datos personales."""

    row: int
    ref: str
    reason: str


@dataclass
class EntityReport:
    created: int = 0
    existing: int = 0
    skipped: list[Issue] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)
    to_review: list[Issue] = field(default_factory=list)


@dataclass
class ImportReport:
    applied: bool
    clients: EntityReport = field(default_factory=EntityReport)
    equipment: EntityReport = field(default_factory=EntityReport)


def read_rows(path: Path, headers: tuple[str, ...]) -> list[tuple[int, Row]]:
    """Filas no vacías de la primera hoja, con su número de fila de la planilla."""
    try:
        # Sin read_only: las exportaciones de Persat no declaran el tamaño de la hoja.
        sheet = load_workbook(path).worksheets[0]
    except (OSError, ValueError, KeyError) as exc:
        raise PersatFileError(f"No se pudo leer {path.name}: {exc}") from None
    rows = sheet.iter_rows(values_only=True)
    header = [persat.text(value) for value in next(rows, ())]
    missing = [name for name in headers if name not in header]
    if missing:
        raise PersatFileError(f"A {path.name} le faltan las columnas: {', '.join(missing)}.")
    result = []
    for number, values in enumerate(rows, start=2):
        cells = {name: persat.text(value) for name, value in zip(header, values, strict=False)}
        if any(cells.values()):
            result.append((number, cells))
    return result


async def import_persat(
    session: AsyncSession,
    *,
    clients: list[tuple[int, Row]] | None,
    equipment: list[tuple[int, Row]] | None,
    apply: bool,
) -> ImportReport:
    report = ImportReport(applied=apply)
    try:
        if clients is not None:
            await _import_clients(session, clients, report.clients)
        if equipment is not None:
            await _import_equipment(session, equipment, report.equipment)
    except BaseException:
        await session.rollback()
        raise
    if apply:
        await session.commit()
    else:
        await session.rollback()
    return report


# ---------- Clientes ----------


async def _import_clients(
    session: AsyncSession, rows: list[tuple[int, Row]], report: EntityReport
) -> None:
    clients = ClientRepository(session)
    seen: set[str] = set()
    without_number: list[tuple[int, Row, str]] = []
    for row, cells in rows:
        number = persat.client_number(cells["Nro. de Cliente"])
        ref = f"cliente {number}" if number else "cliente sin número"
        name = cells["Razón Social"]
        if persat.is_internal_client(cells["Tipo"]):
            report.skipped.append(Issue(row, ref, f"registro interno ({cells['Tipo']})"))
            continue
        if persat.is_test_record(name):
            report.skipped.append(Issue(row, ref, "registro de prueba"))
            continue
        problem = _client_problem(number, cells)
        if problem:
            report.errors.append(Issue(row, ref, problem))
            continue
        legacy_id = number or f"sin-numero:{name}"
        if legacy_id in seen:
            report.errors.append(Issue(row, ref, "repetido en el archivo"))
            continue
        seen.add(legacy_id)
        if await _has_reference(session, CLIENT_ENTITY, legacy_id):
            report.existing += 1
            continue
        if not number:
            # Al final, para que el número nuevo siga a los que vienen de Persat.
            without_number.append((row, cells, legacy_id))
            continue
        if await clients.get_by_number(number) is not None:
            report.errors.append(Issue(row, ref, "ya hay un cliente con ese número cargado a mano"))
            continue
        await _create_client(session, row, cells, number, legacy_id, report)

    for row, cells, legacy_id in without_number:
        number = str((await clients.max_numeric_number() or 0) + 1)
        report.to_review.append(
            Issue(row, f"cliente {number}", "no tenía número en Persat: se le asignó uno nuevo")
        )
        await _create_client(session, row, cells, number, legacy_id, report)


def _client_problem(number: str, cells: Row) -> str | None:
    if not cells["Razón Social"]:
        return "sin razón social"
    description = persat.split_description(cells["Descripción"])
    return _too_long(
        número=number,
        **{"razón social": cells["Razón Social"]},
        dirección=cells["Dirección"],
        referencia=description.reference or "",
    )


async def _create_client(
    session: AsyncSession,
    row: int,
    cells: Row,
    number: str,
    legacy_id: str,
    report: EntityReport,
) -> None:
    description = persat.split_description(cells["Descripción"])
    client = Client(
        id=uuid.uuid4(),
        number=number,
        name=cells["Razón Social"],
        segment=ClientSegment(persat.segment(cells["Tipo"])),
        tax_id=description.tax_id,
        tax_condition=TaxCondition(description.tax_condition)
        if description.tax_condition
        else None,
    )
    address = cells["Dirección"]
    if not address:
        report.to_review.append(Issue(row, f"cliente {number}", "sin dirección"))
    site = ClientSite(
        id=uuid.uuid4(),
        client_id=client.id,
        address=address or NO_ADDRESS,
        reference=description.reference,
    )
    session.add_all([client, site, _reference(CLIENT_ENTITY, client.id, legacy_id)])
    contact = ClientContact(
        client_id=client.id,
        name=cells["Nombre"] or None,
        phone=persat.phone(cells["Telefono"]),
        email=persat.email(cells["Email"]),
    )
    if contact.name or contact.phone or contact.email:
        session.add(contact)
    await session.flush()
    report.created += 1


# ---------- Equipos ----------


async def _import_equipment(
    session: AsyncSession, rows: list[tuple[int, Row]], report: EntityReport
) -> None:
    types = await _equipment_types(session)
    repo = EquipmentRepository(session)
    code_counts = Counter(cells["Numero de Serie"] for _, cells in rows)
    occurrences: Counter[str] = Counter()
    # Los que necesitan código nuevo se crean al final, cuando los de Persat ya ocupan el suyo.
    pending: list[tuple[int, Equipment, Client, EquipmentType, str]] = []

    for row, cells in rows:
        code = cells["Numero de Serie"]
        number = persat.client_number(cells["Nro. de Cliente"])
        ref = f"equipo {code or 'sin código'} (cliente {number or 'sin número'})"
        base_id = f"{number}/{code}"
        occurrences[base_id] += 1
        legacy_id = base_id if occurrences[base_id] == 1 else f"{base_id}#{occurrences[base_id]}"
        problem = _too_long(código=code, ubicación=cells["UBICACION"])
        if problem:
            report.errors.append(Issue(row, ref, problem))
            continue
        if await _has_reference(session, EQUIPMENT_ENTITY, legacy_id):
            report.existing += 1
            continue
        client = await ClientRepository(session).get_by_number(number) if number else None
        if client is None or client.deleted_at is not None:
            report.errors.append(Issue(row, ref, "no se puede asociar a un cliente importado"))
            continue
        site = await _first_active_site(session, client.id)
        if site is None:
            report.errors.append(Issue(row, ref, "el cliente no tiene un domicilio activo"))
            continue

        data = persat.equipment_data(
            raw_type=cells["TIPO"],
            raw_subtype=cells["SUB TIPO"],
            raw_model=cells["MODELO"],
            raw_brand=cells["MARCA"],
            raw_capacity=cells["POTENCIA"],
            raw_gas=cells["GAS REFRIGERANTE"],
            raw_voltage=cells["TENSION DE ALIMENTACION"],
        )
        review = list(data.to_review)
        location = cells["UBICACION"]
        if not location:
            review.append("sin ubicación")
        if not any(cells[name] for name in EQUIPMENT_HEADERS[2:]):
            review.append("sin ningún dato")
        equipment_type = types[data.type_name]
        equipment = Equipment(
            id=uuid.uuid4(),
            client_site_id=site.id,
            equipment_type_id=equipment_type.id,
            internal_code=code,
            location_note=location or persat.NO_LOCATION,
            brand=data.brand,
            notes=data.notes,
        )
        needs_new_code = not code or code_counts[code] > 1 or await repo.code_exists(code)
        for reason in review:
            report.to_review.append(Issue(row, ref, reason))
        if needs_new_code:
            pending.append((row, equipment, client, equipment_type, legacy_id))
            continue
        session.add_all([equipment, _reference(EQUIPMENT_ENTITY, equipment.id, legacy_id)])
        await session.flush()
        report.created += 1

    for row, equipment, client, equipment_type, legacy_id in pending:
        persat_code = equipment.internal_code
        equipment.internal_code = await next_code(session, client, equipment_type)
        if persat_code:
            note = f"Código de Persat: {persat_code}."
            equipment.notes = f"{equipment.notes} {note}" if equipment.notes else note
            reason = f"código repetido en Persat: se le asignó {equipment.internal_code}"
        else:
            reason = f"sin código en Persat: se le asignó {equipment.internal_code}"
        report.to_review.append(Issue(row, f"equipo {persat_code or 'sin código'}", reason))
        session.add_all([equipment, _reference(EQUIPMENT_ENTITY, equipment.id, legacy_id)])
        await session.flush()
        report.created += 1


async def _equipment_types(session: AsyncSession) -> dict[str, EquipmentType]:
    result = await session.execute(select(EquipmentType).where(EquipmentType.deleted_at.is_(None)))
    types = {item.name: item for item in result.scalars()}
    needed = (persat.AIR_CONDITIONER, persat.COMMERCIAL_FRIDGE, persat.HOME_FRIDGE)
    missing = [name for name in needed if name not in types]
    if missing:
        raise PersatFileError(f"Faltan los tipos de equipo: {', '.join(missing)}.")
    return types


async def _first_active_site(session: AsyncSession, client_id: uuid.UUID) -> ClientSite | None:
    sites = await ClientSiteRepository(session).list_for_client(client_id)
    return next((site for site in sites if site.is_active), None)


# ---------- Comunes ----------


async def _has_reference(session: AsyncSession, entity: str, legacy_id: str) -> bool:
    stmt = select(LegacyReference.id).where(
        LegacyReference.entity == entity,
        LegacyReference.source == SOURCE,
        LegacyReference.legacy_id == legacy_id,
        LegacyReference.deleted_at.is_(None),
    )
    return await session.scalar(select(stmt.exists())) or False


def _reference(entity: str, entity_id: uuid.UUID, legacy_id: str) -> LegacyReference:
    return LegacyReference(entity=entity, entity_id=entity_id, source=SOURCE, legacy_id=legacy_id)


def _too_long(**values: str) -> str | None:
    for name, value in values.items():
        if len(value) > _MAX_LENGTHS[name]:
            return f"{name}: más de {_MAX_LENGTHS[name]} caracteres"
    return None
